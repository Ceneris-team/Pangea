"""Copia el estado de los casos y de los defectos a la RTM y deja calculada la hoja Historias.

La RTM no lee los otros libros con fórmulas (se rompen al mover carpetas): este script copia
ID, historia y estado a las hojas «Fuente casos» y «Fuente defectos», y la hoja Historias
cuenta desde ahí. Ejecutar cada vez que cambie un estado en Casos de Prueba o en Defectos:

    backend/.venv/Scripts/python.exe QA/05_Matriz_Trazabilidad_RTM/sincronizar_rtm.py
"""
import sys
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

QA = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(QA))
from estilo_qa import aplicar_estilo_xlsx, estilo_encabezado  # noqa: E402

VERSION = "v1.2"
RTM = QA / "05_Matriz_Trazabilidad_RTM" / f"Cervelab_Matriz_Trazabilidad_RTM_{VERSION}.xlsx"
CASOS = QA / "02_Casos_de_Prueba" / f"Cervelab_Casos_de_Prueba_{VERSION}.xlsx"
DEFECTOS = QA / "04_Defectos" / f"Cervelab_Registro_de_Defectos_{VERSION}.xlsx"

FILAS_HISTORIAS = range(2, 15)  # filas preparadas en la hoja Historias
FIN = 500                       # última fila que leen las fórmulas en las hojas fuente


def leer(libro, hoja, col_estado, prefijo):
    """Devuelve (ID, Historia, Estado) buscando las columnas por su encabezado."""
    ws = openpyxl.load_workbook(libro, data_only=True)[hoja]
    encabezados = [c.value for c in ws[1]]
    i_id, i_hu, i_estado = (encabezados.index(t) for t in ("ID", "Historia", col_estado))
    filas = []
    for fila in ws.iter_rows(min_row=2, values_only=True):
        id_ = fila[i_id]
        if isinstance(id_, str) and id_.startswith(prefijo):
            filas.append((id_, fila[i_hu], fila[i_estado]))
    return filas


def escribir_fuente(wb, nombre, encabezados, filas):
    if nombre in wb.sheetnames:
        del wb[nombre]
    ws = wb.create_sheet(nombre)
    ws.append(encabezados)
    for celda in ws[1]:
        estilo_encabezado(celda, "rtm")
    for fila in filas:
        ws.append(fila)
    ws.freeze_panes = "A2"
    for col, ancho in zip("ABC", (14, 12, 16)):
        ws.column_dimensions[col].width = ancho
    nota = ws.cell(row=1, column=5, value="Se llena con sincronizar_rtm.py desde el libro de origen. No editar a mano.")
    nota.font = Font(italic=True, color="5F6870")


def formulas_historias(ws):
    casos = "'Fuente casos'!$B$2:$B${0}".format(FIN)
    estado = "'Fuente casos'!$C$2:$C${0}".format(FIN)
    d_id = "'Fuente defectos'!$A$2:$A${0}".format(FIN)
    d_hu = "'Fuente defectos'!$B$2:$B${0}".format(FIN)
    d_est = "'Fuente defectos'!$C$2:$C${0}".format(FIN)
    for r in FILAS_HISTORIAS:
        si = f'=IF($A{r}="","",'
        ws[f"K{r}"] = f"{si}COUNTIF({casos},$A{r}))"
        for col, valor in (("L", "Aprobado"), ("M", "Fallido"), ("N", "Bloqueado")):
            ws[f"{col}{r}"] = f'{si}COUNTIFS({casos},$A{r},{estado},"{valor}"))'
        ws[f"O{r}"] = (f'{si}COUNTIFS({d_hu},$A{r},{d_id},"DEF-*",'
                       f'{d_est},"<>Cerrado",{d_est},"<>Rechazado"))')
    ws["K1"] = "Casos de la HU"


def portada(ws):
    ws["B10"] = "Scrum Master: Jareth Wilver Trujillo  /  Product Owner: Fabio Saavedra"
    ws["A19"] = "Fuente casos y Fuente defectos"
    ws["B19"] = ("Copia del ID, la historia y el estado de cada caso y defecto. La hoja Historias cuenta desde ahí. "
                 "Se actualiza con sincronizar_rtm.py; no se edita a mano.")
    for col in "AB":
        ws[f"{col}19"]._style = copy(ws[f"{col}18"]._style)


def main():
    casos = leer(CASOS, "Casos de Prueba", "Estado", "CP-")
    defectos = leer(DEFECTOS, "Defectos", "Estado", "DEF-")
    wb = openpyxl.load_workbook(RTM)
    escribir_fuente(wb, "Fuente casos", ["ID", "Historia", "Estado"], casos)
    escribir_fuente(wb, "Fuente defectos", ["ID", "Historia", "Estado"], defectos)
    formulas_historias(wb["Historias"])
    portada(wb["Portada"])
    aplicar_estilo_xlsx(wb, "rtm")
    wb.calculation.fullCalcOnLoad = True
    wb.save(RTM)
    print(f"RTM sincronizada: {len(casos)} casos, {len(defectos)} defectos")


if __name__ == "__main__":
    main()
